"""
Synthetic document generator for KYC testing.

Generates realistic PDFs and DOCX files for India (IN) and Singapore (SG)
with controlled variations to test extraction and gap detection.
"""

from __future__ import annotations

import random
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Literal

try:
    from docx import Document as DocxDocument  # type: ignore
except Exception:  # pragma: no cover
    DocxDocument = None  # type: ignore
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


@dataclass
class CompanyProfile:
    name: str
    registration_number: str  # CIN for IN, UEN for SG
    incorporation_date: str
    registered_address: str
    country: Literal["IN", "SG"]
    entity_type: Literal["company", "llp"] = "company"
    pan: str | None = None
    llpin: str | None = None


def random_date_str(days_ago_max: int = 3650) -> str:
    d = datetime.now() - timedelta(days=random.randint(100, days_ago_max))
    return d.strftime("%Y-%m-%d")


def _random_pan() -> str:
    letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    return (
        "".join(random.choice(letters) for _ in range(5))
        + "".join(str(random.randint(0, 9)) for _ in range(4))
        + random.choice(letters)
    )


def _random_llpin() -> str:
    # Common demo-friendly LLPIN format used in practice: AAA-1234
    letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    return f"{random.choice(letters)}{random.choice(letters)}{random.choice(letters)}-{random.randint(1000,9999)}"


def generate_company_profile(country: Literal["IN", "SG"], *, entity_type: Literal["company", "llp"] = "company") -> CompanyProfile:
    if country == "IN":
        if entity_type == "llp":
            names = [
                "TechVenture Consulting LLP",
                "Global Trade Partners LLP",
                "Innovate Manufacturing LLP",
                "Alpha Services LLP",
            ]
            reg = _random_llpin()
        else:
            names = [
                "TechVenture Solutions Private Limited",
                "Global Trade Partners India Pvt Ltd",
                "Innovate Manufacturing Co Ltd",
                "Alpha Services India Limited",
            ]
            # CIN format: L12345AB1234ABC123456 (simplified)
            reg = f"L{random.randint(10000,99999)}{random.choice(['MH','DL','KA','TN'])}{random.randint(1000,9999)}PLC{random.randint(100000,999999)}"
        addresses = [
            "123 Business Park, Andheri East, Mumbai 400069, Maharashtra",
            "456 Corporate Tower, Connaught Place, New Delhi 110001",
            "789 Tech Hub, Electronic City, Bangalore 560100, Karnataka",
        ]
        return CompanyProfile(
            name=random.choice(names),
            registration_number=reg,
            incorporation_date=random_date_str(),
            registered_address=random.choice(addresses),
            country="IN",
            entity_type=entity_type,
            pan=_random_pan(),
            llpin=reg if entity_type == "llp" else None,
        )
    else:  # SG
        names = [
            "Asia Pacific Trading Pte Ltd",
            "Singapore Fintech Solutions Pte Ltd",
            "Global Shipping Partners Pte Ltd",
            "Marina Bay Consulting Pte Ltd",
        ]
        # UEN format: TYYPQXXXXX
        uen = f"T{random.randint(18,24)}{random.choice(['A','B','C','D','E'])}{random.choice(['A','B','C'])}{random.randint(10000,99999)}"
        addresses = [
            "1 Raffles Place, #12-01 One Raffles Place, Singapore 048616",
            "8 Marina Boulevard, #11-01 Marina Bay Financial Centre, Singapore 018981",
            "3 Temasek Avenue, #15-01 Centennial Tower, Singapore 039190",
        ]
        return CompanyProfile(
            name=random.choice(names),
            registration_number=uen,
            incorporation_date=random_date_str(),
            registered_address=random.choice(addresses),
            country="SG",
            entity_type="company",
        )


def generate_certificate_of_incorporation(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    """Generate Certificate of Incorporation PDF."""
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []

    title_style = ParagraphStyle(
        "Title",
        fontSize=18,
        alignment=1,  # center
        spaceAfter=30,
        textColor=colors.HexColor("#1a365d"),
    )
    heading_style = ParagraphStyle(
        "Heading",
        fontSize=12,
        spaceAfter=6,
        textColor=colors.HexColor("#2d3748"),
    )
    normal_style = ParagraphStyle(
        "Normal",
        fontSize=11,
        spaceAfter=12,
        textColor=colors.black,
    )

    # Title
    story.append(Paragraph("CERTIFICATE OF INCORPORATION", title_style))
    story.append(Spacer(1, 0.5 * cm))

    # Government header simulation
    if company.country == "IN":
        story.append(Paragraph("Government of India<br/>Ministry of Corporate Affairs", heading_style))
    else:
        story.append(Paragraph("Accounting and Corporate Regulatory Authority<br/>Republic of Singapore", heading_style))
    story.append(Spacer(1, 0.5 * cm))

    # Content
    content = []
    if "company_name" not in missing_fields:
        content.append(["Company Name:", company.name])
    else:
        content.append(["Company Name:", "[REDACTED]"])

    if "registration_number" not in missing_fields:
        reg_label = "CIN" if company.country == "IN" else "UEN"
        content.append([f"{reg_label}:", company.registration_number])

    if "incorporation_date" not in missing_fields:
        content.append(["Incorporation Date:", company.incorporation_date])

    if "registered_address" not in missing_fields:
        content.append(["Registered Address:", company.registered_address])

    table = Table(content, colWidths=[6 * cm, 10 * cm])
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#edf2f7")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 11),
            ("GRID", (0, 0), (-1, -1), 1, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ])
    )
    story.append(table)
    story.append(Spacer(1, 1 * cm))

    # Certification text
    story.append(
        Paragraph(
            "This is to certify that the above company is incorporated under the relevant Companies Act.",
            normal_style,
        )
    )

    doc.build(story)


def generate_ubo_declaration(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    """Generate UBO Declaration PDF."""
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []

    title_style = ParagraphStyle(
        "Title",
        fontSize=16,
        alignment=1,
        spaceAfter=20,
        textColor=colors.HexColor("#1a365d"),
    )
    normal_style = ParagraphStyle(
        "Normal",
        fontSize=11,
        spaceAfter=12,
        textColor=colors.black,
    )

    story.append(Paragraph("ULTIMATE BENEFICIAL OWNER DECLARATION", title_style))
    story.append(Spacer(1, 0.5 * cm))

    if "entity_name" not in missing_fields and "company_name" not in missing_fields:
        story.append(Paragraph(f"<b>Entity Name:</b> {company.name}", normal_style))

    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph("<b>Beneficial Owner Details:</b>", normal_style))

    # UBO details
    ubo_first = random.choice(["Rajesh", "Priya", "Suresh", "Wei Ming", "Xiao Li"])
    ubo_last = random.choice(["Kumar", "Sharma", "Patel", "Tan", "Lim"])
    ubo_name = f"{ubo_first} {ubo_last}"

    content = []
    if "ubo_name" not in missing_fields:
        content.append(["UBO Name:", ubo_name])
    else:
        content.append(["UBO Name:", ""])

    if "ubo_dob" not in missing_fields:
        dob = random_date_str(15000)  # Birth date
        content.append(["UBO DOB:", dob])

    if "ubo_nationality" not in missing_fields:
        nationality = "Indian" if company.country == "IN" else random.choice(["Singaporean", "Malaysian"])
        content.append(["UBO Nationality:", nationality])

    if "ownership_percent" not in missing_fields:
        content.append(["Ownership %:", str(random.choice([25, 30, 40, 51, 75]))])

    table = Table(content, colWidths=[5 * cm, 11 * cm])
    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#edf2f7")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 11),
            ("GRID", (0, 0), (-1, -1), 1, colors.grey),
        ])
    )
    story.append(table)

    story.append(Spacer(1, 1 * cm))
    story.append(
        Paragraph(
            "I declare that the above information is true and correct to the best of my knowledge.",
            normal_style,
        )
    )

    doc.build(story)


def generate_proof_of_address(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    """Generate utility bill as proof of address PDF."""
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []

    normal_style = ParagraphStyle(
        "Normal",
        fontSize=10,
        spaceAfter=8,
        textColor=colors.black,
    )
    heading_style = ParagraphStyle(
        "Heading",
        fontSize=14,
        spaceAfter=12,
        textColor=colors.HexColor("#2d3748"),
    )

    # Utility header
    utilities = ["Electricity Board", "Water Supply Board", "Telecom Services"]
    utility = random.choice(utilities)
    story.append(Paragraph(f"<b>{utility}</b>", heading_style))
    story.append(Spacer(1, 0.3 * cm))

    # Bill details
    bill_date = (datetime.now() - timedelta(days=random.randint(10, 60))).strftime("%Y-%m-%d")
    story.append(Paragraph(f"Bill Date: {bill_date}", normal_style))
    story.append(Spacer(1, 0.2 * cm))

    if "company_name" not in missing_fields:
        story.append(Paragraph(f"Customer: {company.name}", normal_style))

    story.append(Spacer(1, 0.3 * cm))

    if "registered_address" not in missing_fields:
        story.append(Paragraph(f"Service Address:<br/>{company.registered_address}", normal_style))
    else:
        story.append(Paragraph("Service Address:<br/>[Address details missing]", normal_style))

    story.append(Spacer(1, 0.5 * cm))

    # Add some bill content
    bill_amount = random.randint(1000, 50000)
    story.append(Paragraph(f"Bill Amount: ₹{bill_amount:,}" if company.country == "IN" else f"Bill Amount: S${bill_amount:,}", normal_style))
    story.append(Paragraph(f"Due Date: {(datetime.now() + timedelta(days=15)).strftime('%Y-%m-%d')}", normal_style))

    doc.build(story)


def generate_directors_list_docx(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    """Generate Directors List as DOCX."""
    if DocxDocument is None:
        raise RuntimeError("python-docx is not installed. Install backend requirements or use PDF outputs.")
    missing_fields = missing_fields or set()
    doc = DocxDocument()

    # Title
    title = doc.add_heading("Company Profile and Directors List", 0)
    title.alignment = 1  # center

    if "company_name" not in missing_fields:
        doc.add_paragraph(f"Company Name: {company.name}")

    doc.add_paragraph(f"Document Date: {datetime.now().strftime('%Y-%m-%d')}")
    doc.add_paragraph()

    # Directors table
    doc.add_heading("Board of Directors", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Name"
    hdr_cells[1].text = "Designation"
    hdr_cells[2].text = "Date Appointed"

    directors = [
        ("Rajesh Kumar", "Managing Director", "2018-03-15"),
        ("Priya Sharma", "Executive Director", "2019-07-20"),
        ("Amit Patel", "Non-Executive Director", "2020-01-10"),
    ]

    for name, desig, date in directors:
        row_cells = table.add_row().cells
        row_cells[0].text = name
        row_cells[1].text = desig
        row_cells[2].text = date

    if "directors" in missing_fields:
        # Remove the table content to simulate missing info
        for row in table.rows[1:]:
            for cell in row.cells:
                cell.text = ""

    doc.add_paragraph()
    doc.add_paragraph("This document certifies the above as current directors of the company.")

    doc.save(str(output_path))


def generate_directors_list_pdf(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []

    title_style = ParagraphStyle(
        "Title",
        fontSize=16,
        alignment=1,
        spaceAfter=18,
        textColor=colors.HexColor("#1a365d"),
    )
    normal_style = ParagraphStyle(
        "Normal",
        fontSize=11,
        spaceAfter=10,
        textColor=colors.black,
    )

    story.append(Paragraph("DIRECTORS LIST / COMPANY PROFILE", title_style))
    if "company_name" not in missing_fields:
        story.append(Paragraph(f"<b>Company Name:</b> {company.name}", normal_style))
    story.append(Paragraph(f"<b>Document Date:</b> {datetime.now().strftime('%Y-%m-%d')}", normal_style))
    story.append(Spacer(1, 0.4 * cm))

    directors = [
        ("Rajesh Kumar", "Managing Director", "2018-03-15"),
        ("Priya Sharma", "Executive Director", "2019-07-20"),
        ("Amit Patel", "Non-Executive Director", "2020-01-10"),
    ]
    rows = [["Name", "Designation", "Date Appointed"]]
    if "directors" not in missing_fields:
        rows.extend([[a, b, c] for (a, b, c) in directors])
    else:
        rows.extend([["", "", ""] for _ in directors])

    t = Table(rows, colWidths=[6 * cm, 6 * cm, 4 * cm])
    t.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf2f7")),
            ("GRID", (0, 0), (-1, -1), 1, colors.grey),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
        ])
    )
    story.append(t)
    story.append(Spacer(1, 0.6 * cm))
    story.append(Paragraph("This document certifies the above as current directors of the company.", normal_style))
    doc.build(story)


def generate_pan_card(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []
    title_style = ParagraphStyle("Title", fontSize=16, alignment=1, spaceAfter=16, textColor=colors.HexColor("#1a365d"))
    normal_style = ParagraphStyle("Normal", fontSize=11, spaceAfter=10, textColor=colors.black)

    story.append(Paragraph("INCOME TAX DEPARTMENT, GOVT. OF INDIA", title_style))
    story.append(Paragraph("ENTITY PAN DETAILS", ParagraphStyle("Sub", fontSize=12, alignment=1, spaceAfter=18)))

    rows = []
    if "entity_name" not in missing_fields:
        rows.append(["Entity Name:", company.name])
    if "pan" not in missing_fields:
        rows.append(["PAN:", company.pan or _random_pan()])
    t = Table(rows, colWidths=[5 * cm, 11 * cm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#edf2f7")), ("GRID", (0, 0), (-1, -1), 1, colors.grey)]))
    story.append(t)
    story.append(Spacer(1, 0.8 * cm))
    story.append(Paragraph("This is a synthetic document generated for KYC testing.", normal_style))
    doc.build(story)


def generate_moa_aoa(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []
    title_style = ParagraphStyle("Title", fontSize=16, alignment=1, spaceAfter=18, textColor=colors.HexColor("#1a365d"))
    normal_style = ParagraphStyle("Normal", fontSize=11, spaceAfter=10, textColor=colors.black, leading=14)

    story.append(Paragraph("MEMORANDUM & ARTICLES OF ASSOCIATION (EXTRACT)", title_style))
    if "company_name" not in missing_fields:
        story.append(Paragraph(f"<b>Company Name:</b> {company.name}", normal_style))
    story.append(Spacer(1, 0.3 * cm))
    if "registered_office" not in missing_fields:
        story.append(Paragraph(f"<b>Registered Office Address:</b><br/>{company.registered_address}", normal_style))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph("<b>Objects Clause (Summary):</b> The principal objects of the company include software services, consulting, and related activities.", normal_style))
    doc.build(story)


def generate_llp_incorporation_certificate(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []
    title_style = ParagraphStyle("Title", fontSize=16, alignment=1, spaceAfter=20, textColor=colors.HexColor("#1a365d"))
    normal_style = ParagraphStyle("Normal", fontSize=11, spaceAfter=12, textColor=colors.black)

    story.append(Paragraph("CERTIFICATE OF INCORPORATION (LLP)", title_style))
    story.append(Paragraph("Government of India<br/>Ministry of Corporate Affairs", ParagraphStyle("H", fontSize=12, spaceAfter=14)))

    rows = []
    if "llp_name" not in missing_fields:
        rows.append(["LLP Name:", company.name])
    if "llpin" not in missing_fields:
        rows.append(["LLPIN:", company.llpin or company.registration_number])
    if "incorporation_date" not in missing_fields:
        rows.append(["Incorporation Date:", company.incorporation_date])
    if "registered_address" not in missing_fields:
        rows.append(["Registered Address:", company.registered_address])
    t = Table(rows, colWidths=[5 * cm, 11 * cm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#edf2f7")), ("GRID", (0, 0), (-1, -1), 1, colors.grey)]))
    story.append(t)
    story.append(Spacer(1, 0.8 * cm))
    story.append(Paragraph("This is to certify that the above LLP is incorporated under the LLP Act.", normal_style))
    doc.build(story)


def generate_llp_agreement(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []
    title_style = ParagraphStyle("Title", fontSize=16, alignment=1, spaceAfter=18, textColor=colors.HexColor("#1a365d"))
    normal_style = ParagraphStyle("Normal", fontSize=11, spaceAfter=10, textColor=colors.black, leading=14)

    story.append(Paragraph("LLP AGREEMENT (EXTRACT)", title_style))
    if "llp_name" not in missing_fields:
        story.append(Paragraph(f"<b>LLP Name:</b> {company.name}", normal_style))
    eff = (datetime.now() - timedelta(days=random.randint(30, 365))).strftime("%Y-%m-%d")
    if "effective_date" not in missing_fields:
        story.append(Paragraph(f"<b>Effective Date:</b> {eff}", normal_style))
    if "partners" not in missing_fields:
        story.append(Paragraph("<b>Partners:</b> Rajesh Kumar; Priya Sharma; Amit Patel", normal_style))
    story.append(Paragraph("This is a synthetic extract for KYC testing only.", normal_style))
    doc.build(story)


def generate_authorization_letter(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []
    title_style = ParagraphStyle("Title", fontSize=16, alignment=1, spaceAfter=18, textColor=colors.HexColor("#1a365d"))
    normal_style = ParagraphStyle("Normal", fontSize=11, spaceAfter=10, textColor=colors.black, leading=14)

    story.append(Paragraph("AUTHORIZATION LETTER (ACCOUNT OPENING)", title_style))
    dt = (datetime.now() - timedelta(days=random.randint(1, 30))).strftime("%Y-%m-%d")
    if "authorization_date" not in missing_fields:
        story.append(Paragraph(f"<b>Date:</b> {dt}", normal_style))
    if "entity_name" not in missing_fields:
        story.append(Paragraph(f"<b>Entity Name:</b> {company.name}", normal_style))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph("We hereby authorize the following signatories to open and operate bank account(s) on behalf of the entity:", normal_style))
    if "authorized_signatories" not in missing_fields:
        story.append(Paragraph("<b>Authorized Signatories:</b> Rajesh Kumar; Priya Sharma", normal_style))
    doc.build(story)


def generate_partner_list(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []
    title_style = ParagraphStyle("Title", fontSize=16, alignment=1, spaceAfter=18, textColor=colors.HexColor("#1a365d"))
    normal_style = ParagraphStyle("Normal", fontSize=11, spaceAfter=10, textColor=colors.black)

    story.append(Paragraph("PARTNERS LIST", title_style))
    if "firm_name" not in missing_fields and "llp_name" not in missing_fields and "entity_name" not in missing_fields:
        story.append(Paragraph(f"<b>Entity Name:</b> {company.name}", normal_style))
    story.append(Spacer(1, 0.3 * cm))
    rows = [["Name", "Role"]]
    if "partners" not in missing_fields:
        rows.extend([["Rajesh Kumar", "Designated Partner"], ["Priya Sharma", "Designated Partner"], ["Amit Patel", "Partner"]])
    else:
        rows.extend([["", ""], ["", ""], ["", ""]])
    t = Table(rows, colWidths=[10 * cm, 6 * cm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf2f7")), ("GRID", (0, 0), (-1, -1), 1, colors.grey)]))
    story.append(t)
    doc.build(story)


def generate_board_resolution(
    company: CompanyProfile,
    output_path: Path,
    *,
    missing_fields: set[str] | None = None,
) -> None:
    """Generate Board Resolution PDF (for SG primarily)."""
    missing_fields = missing_fields or set()
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []

    title_style = ParagraphStyle(
        "Title",
        fontSize=16,
        alignment=1,
        spaceAfter=20,
        textColor=colors.HexColor("#1a365d"),
    )
    normal_style = ParagraphStyle(
        "Normal",
        fontSize=11,
        spaceAfter=12,
        textColor=colors.black,
        leading=14,
    )

    story.append(Paragraph("BOARD RESOLUTION", title_style))
    story.append(Spacer(1, 0.5 * cm))

    res_date = (datetime.now() - timedelta(days=random.randint(5, 60))).strftime("%Y-%m-%d")

    if "company_name" not in missing_fields:
        story.append(Paragraph(f"<b>Company:</b> {company.name}", normal_style))

    if "resolution_date" not in missing_fields:
        story.append(Paragraph(f"<b>Date of Resolution:</b> {res_date}", normal_style))

    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph(
        "<b>RESOLVED THAT:</b> The Company hereby authorizes the opening of bank account(s) with the Bank and "
        "authorizes the following persons to operate the said account(s):",
        normal_style,
    ))

    story.append(Spacer(1, 0.3 * cm))

    signatories = [
        ("Rajesh Kumar", "Managing Director"),
        ("Priya Sharma", "Executive Director"),
    ]

    for i, (name, title) in enumerate(signatories, 1):
        story.append(Paragraph(f"{i}. {name} - {title}", normal_style))

    story.append(Spacer(1, 1 * cm))
    story.append(Paragraph("This resolution was passed at a duly convened meeting of the Board of Directors.", normal_style))

    doc.build(story)


def generate_all_documents(
    output_dir: Path,
    country: Literal["IN", "SG"],
    variations: list[dict] | None = None,
) -> list[Path]:
    """Generate all document types for a country.
    
    Args:
        output_dir: Base output directory
        country: Country code
        variations: List of variation configs for generating multiple sets
    
    Returns:
        List of generated file paths
    """
    output_dir = Path(output_dir)
    country_dir = output_dir / country
    country_dir.mkdir(parents=True, exist_ok=True)

    generated: list[Path] = []

    if variations is None:
        variations = [{}]

    for i, var in enumerate(variations):
        suffix = f"_set{i+1}" if len(variations) > 1 else ""
        missing_fields = var.get("missing_fields", set())

        if country == "IN":
            # Generate both Company and LLP packs (PDFs) for demo completeness.
            for entity in ("company", "llp"):
                entity_dir = country_dir / entity
                entity_dir.mkdir(parents=True, exist_ok=True)
                profile = generate_company_profile(country, entity_type=entity)  # type: ignore[arg-type]

                if entity == "company":
                    coi_path = entity_dir / f"certificate_of_incorporation{suffix}.pdf"
                    generate_certificate_of_incorporation(profile, coi_path, missing_fields=missing_fields)
                    generated.append(coi_path)

                    pan_path = entity_dir / f"pan_card{suffix}.pdf"
                    generate_pan_card(profile, pan_path, missing_fields=missing_fields)
                    generated.append(pan_path)

                    moa_path = entity_dir / f"moa_aoa{suffix}.pdf"
                    generate_moa_aoa(profile, moa_path, missing_fields=missing_fields)
                    generated.append(moa_path)

                    dir_path = entity_dir / f"directors_list{suffix}.pdf"
                    generate_directors_list_pdf(profile, dir_path, missing_fields=missing_fields)
                    generated.append(dir_path)

                    br_path = entity_dir / f"board_resolution{suffix}.pdf"
                    generate_board_resolution(profile, br_path, missing_fields=missing_fields)
                    generated.append(br_path)

                    ubo_path = entity_dir / f"ubo_declaration{suffix}.pdf"
                    generate_ubo_declaration(profile, ubo_path, missing_fields=missing_fields)
                    generated.append(ubo_path)

                    poa_path = entity_dir / f"proof_of_address{suffix}.pdf"
                    generate_proof_of_address(profile, poa_path, missing_fields=missing_fields)
                    generated.append(poa_path)

                else:
                    llp_coi = entity_dir / f"llp_incorporation_certificate{suffix}.pdf"
                    generate_llp_incorporation_certificate(profile, llp_coi, missing_fields=missing_fields)
                    generated.append(llp_coi)

                    pan_path = entity_dir / f"pan_card{suffix}.pdf"
                    generate_pan_card(profile, pan_path, missing_fields=missing_fields)
                    generated.append(pan_path)

                    llp_ag = entity_dir / f"llp_agreement{suffix}.pdf"
                    generate_llp_agreement(profile, llp_ag, missing_fields=missing_fields)
                    generated.append(llp_ag)

                    partners = entity_dir / f"partner_list{suffix}.pdf"
                    generate_partner_list(profile, partners, missing_fields=missing_fields)
                    generated.append(partners)

                    auth = entity_dir / f"authorization_letter{suffix}.pdf"
                    generate_authorization_letter(profile, auth, missing_fields=missing_fields)
                    generated.append(auth)

                    ubo_path = entity_dir / f"ubo_declaration{suffix}.pdf"
                    generate_ubo_declaration(profile, ubo_path, missing_fields=missing_fields)
                    generated.append(ubo_path)

                    poa_path = entity_dir / f"proof_of_address{suffix}.pdf"
                    generate_proof_of_address(profile, poa_path, missing_fields=missing_fields)
                    generated.append(poa_path)

            continue

        else:
            company = generate_company_profile(country)

            # Certificate of Incorporation
            coi_path = country_dir / f"certificate_of_incorporation{suffix}.pdf"
            generate_certificate_of_incorporation(company, coi_path, missing_fields=missing_fields)
            generated.append(coi_path)

            # UBO Declaration
            ubo_path = country_dir / f"ubo_declaration{suffix}.pdf"
            generate_ubo_declaration(company, ubo_path, missing_fields=missing_fields)
            generated.append(ubo_path)

            # Proof of Address
            poa_path = country_dir / f"proof_of_address{suffix}.pdf"
            generate_proof_of_address(company, poa_path, missing_fields=missing_fields)
            generated.append(poa_path)

            # SG uses Board Resolution instead
            br_path = country_dir / f"board_resolution{suffix}.pdf"
            generate_board_resolution(
                company, br_path, missing_fields=missing_fields
            )
            generated.append(br_path)

            # Directors list as PDF for better demo consistency
            dir_path = country_dir / f"directors_list{suffix}.pdf"
            generate_directors_list_pdf(company, dir_path, missing_fields=missing_fields)
            generated.append(dir_path)

    return generated


def main():
    """CLI entry point for generating synthetic documents."""
    import argparse

    parser = argparse.ArgumentParser(description="Generate synthetic KYC documents")
    parser.add_argument("--output", "-o", type=Path, default=Path("out"), help="Output directory")
    parser.add_argument("--country", "-c", choices=["IN", "SG", "both"], default="both")
    parser.add_argument("--sets", "-n", type=int, default=1, help="Number of document sets per country")
    parser.add_argument("--missing", "-m", nargs="+", help="Fields to intentionally omit for testing gaps")

    args = parser.parse_args()

    # Create output directory
    args.output.mkdir(parents=True, exist_ok=True)

    countries = ["IN", "SG"] if args.country == "both" else [args.country]

    for country in countries:
        print(f"\nGenerating documents for {country}...")
        
        variations = []
        for i in range(args.sets):
            var = {}
            if args.missing and i == 0:  # Only omit in first set if testing gaps
                var["missing_fields"] = set(args.missing)
            variations.append(var)

        paths = generate_all_documents(args.output, country, variations)
        for p in paths:
            print(f"  Created: {p}")

    print(f"\nDone! Documents saved to: {args.output.absolute()}")


if __name__ == "__main__":
    main()
